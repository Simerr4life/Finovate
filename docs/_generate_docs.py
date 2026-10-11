#!/usr/bin/env python3
"""Generate Finovate Markdown docs, DOCX, and PPTX deliverables."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGB
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches as PInches, Pt as PPt

ROOT = Path("/workspace")
DOCS = ROOT / "docs"
RAW = DOCS / "assets" / "screenshots" / "raw"
ANN = DOCS / "assets" / "screenshots" / "annotated"
MOB = DOCS / "assets" / "screenshots" / "mobile"
DELIV = DOCS / "deliverables"
ART = Path("/opt/cursor/artifacts")
DELIV.mkdir(parents=True, exist_ok=True)

TODAY = "2026-10-11"
VERSION = "1.0.0"
NAVY = RGBColor(0x0F, 0x2A, 0x22)
GREEN = RGBColor(0x1F, 0x6B, 0x4A)


def write(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content.rstrip() + "\n", encoding="utf-8")
    print("md", path.relative_to(ROOT))


def md_img(rel_from_docs: str, alt: str) -> str:
    if (DOCS / rel_from_docs).exists():
        return f"![{alt}]({rel_from_docs})"
    return f"*Screenshot not available: `{rel_from_docs}`*"


def generate_markdown() -> None:
    write(
        DOCS / "README.md",
        f"""# Finovate Documentation

**Version:** {VERSION}
**Last reviewed:** {TODAY}
**Audience:** End users, developers, administrators, and contributors

Finovate is a personal-finance web application for tracking spending, planning budgets, working toward savings goals, reviewing a personalized savings plan, and previewing **simulated** round-up / ETF investing. Identity is provided by Auth0 and/or Supabase Auth when configured. This documentation describes **verified implementation behavior** from the source code in this repository.

## How to use this documentation

| Role | Start here |
|---|---|
| New end user | [Getting started](getting-started.md) → [User guide](user-guide/README.md) |
| Developer | [Overview](overview.md) → [Architecture](architecture/README.md) → [Development](development/README.md) |
| Admin / deployer | [Deployment](deployment/README.md) → [Security](security/README.md) → [Integrations](integrations/README.md) |
| QA | [Testing](testing/README.md) → [Coverage matrix](coverage-matrix.md) |

## Contents

1. [Product overview](overview.md)
2. [Getting started](getting-started.md)
3. [Feature inventory](feature-inventory.md)
4. [User guide](user-guide/README.md)
5. [Features](features/README.md)
6. [Workflows](workflows/README.md)
7. [Authentication](authentication/README.md)
8. [Account management](account-management/README.md)
9. [Interface](interface/README.md)
10. [API](api/README.md)
11. [Database](database/README.md)
12. [Architecture](architecture/README.md)
13. [Security](security/README.md)
14. [Integrations](integrations/README.md)
15. [Troubleshooting](troubleshooting.md)
16. [Development](development/README.md)
17. [Deployment](deployment/README.md)
18. [Testing](testing/README.md)
19. [Glossary](glossary.md)
20. [Coverage matrix](coverage-matrix.md)
21. [Audit report](audit-report.md)
22. [Changelog](CHANGELOG.md)

## Packaged deliverables

| Format | Path |
|---|---|
| Word document | [deliverables/Finovate-Complete-Documentation.docx](deliverables/Finovate-Complete-Documentation.docx) |
| PowerPoint | [deliverables/Finovate-Visual-Feature-Guide.pptx](deliverables/Finovate-Visual-Feature-Guide.pptx) |

## Screenshots

Real application screenshots (synthetic `demo` data) under `assets/screenshots/`:

- `raw/` — unmodified captures
- `annotated/` — numbered callouts
- `mobile/` — mobile viewport captures

## Honest boundaries

- Investing is **simulated**; no brokerage is connected.
- Bank connection is a **placeholder** (“coming soon”).
- Notification events are **recorded**, not delivered externally.
- Django does **not** send iMessage; Spectrum inbound webhooks are verified.
- Website login requires Auth0 and/or Supabase Auth; `seed_demo` is a fixture for commands and local verification, not a public password login when those providers are expected.
""",
    )

    write(
        DOCS / "overview.md",
        f"""# Product overview

**Last reviewed:** {TODAY}

## Purpose

Finovate helps individuals understand recorded cash flow, plan a budget, save toward goals, receive a **recommendation-only** savings plan, and explore a **simulated** investing experience — without requiring a live bank or brokerage connection.

## Capabilities (verified)

| Area | Capability | Real vs simulated |
|---|---|---|
| Transactions | Manual CRUD, search, CSV export | Real (local DB) |
| Imports | CSV bank-statement upload → preview → commit | Real |
| Statements | Demo/fictional ledger + reconciliation | Demo data |
| Budgets | Periods, category limits, progress, income/expenses | Real math on recorded data |
| Savings goals | Goals, contributions, milestones, pace | Real |
| Savings plan | Scenario engine, allocations, debts | Recommendation only |
| Investing | Round-ups, surplus propose/approve, demo AI, market explore | Simulated |
| Insights / dashboard | Charts, MoM, category breakdown | Derived from recorded data |
| Auth | Auth0 OIDC + optional Supabase email/password | External IdP |
| Messaging assistant | Keyword intents via Photon Spectrum webhook | Inbound verified; outbound not from Django |

## System overview

```mermaid
flowchart LR
  Browser[Browser / Django templates]
  Django[Django apps]
  DB[(SQLite or Postgres)]
  Auth0[Auth0]
  Supa[Supabase Auth]
  TD[Twelve Data]
  Spec[Photon Spectrum]

  Browser --> Django
  Django --> DB
  Browser --> Auth0
  Auth0 --> Django
  Browser --> Supa
  Supa --> Django
  Django --> TD
  Spec --> Django
```

## Apps

| Django app | Owns |
|---|---|
| `core` | Landing, dashboard, insights, money helpers, seeds |
| `accounts` | Profile, Auth0/Supabase, onboarding, settings, `/api/me/` |
| `transactions` | Account, Category, Transaction, statements |
| `budgets` | Budget periods/limits, income, recurring expenses, txn UI |
| `savings` | Goals, contributions, pools, plan engine, debts |
| `investing` | Round-ups, simulated orders, market data, Twelve Data |
| `imports` | CSV import batches |
| `assistant` | Spectrum webhook, intents, link codes |

## Intended users

- Individuals tracking personal finances locally or on hosted Postgres
- Developers extending Finovate
- Operators deploying to Vercel with Auth0/Supabase

## Non-goals (current product)

- Real brokerage order routing
- Live bank aggregation
- Production notification delivery (email/push/SMS)
- Outbound iMessage from the Django process
""",
    )

    write(
        DOCS / "getting-started.md",
        f"""# Getting started

**Last reviewed:** {TODAY}

## Prerequisites

- Python 3.13 recommended (3.12 used successfully for local docs capture)
- Ability to install from `requirements.txt`
- Optional: Auth0, Supabase, Twelve Data, Photon Spectrum credentials

## Local setup

```bash
cd Finovate
python -m venv .venv
source .venv/bin/activate
pip install -r requirements.txt
python manage.py migrate
python manage.py seed_demo
python manage.py runserver
```

Open `http://127.0.0.1:8000/`.

### Optional seeds

```bash
python manage.py seed_peachtree_demo
python manage.py seed_excel_demo_market
python manage.py seed_market_data
```

### Admin

```bash
python manage.py createsuperuser
```

Visit `/admin/`.

## Configuration

Copy `.env.example` → `.env`. SQLite demo data works without `.env`, but **website signup/login stay unavailable until Auth0 and/or Supabase Auth are configured**.

| Variable | Purpose |
|---|---|
| `FINOVATE_SECRET_KEY` | Django secret |
| `AUTH0_*` | Auth0 web app + session encryption |
| `APP_BASE_URL` | Origin for Auth0 callback `/callback` |
| `SUPABASE_URL` + keys | Email/password auth (+ optional Postgres) |
| `DATABASE_URL` / `SUPABASE_DB_URL` | Postgres instead of SQLite |
| `TWELVE_DATA_API_KEY` | Live demo quotes/history |
| `IMESSAGE_*` / `PHOTON_WEBHOOK_SECRET` | Spectrum webhook verification |

Never commit real secret values.

## First-user journey (auth configured)

1. Landing → **Sign up**
2. Auth0 Universal Login or Supabase email form
3. Email verification gate if needed (`/accounts/verify/`)
4. Onboarding: CSV import, manual setup, or continue
5. Preferences → marks `onboarding_completed`
6. Dashboard

## Quick test

```bash
python manage.py test
```
""",
    )

    write(
        DOCS / "feature-inventory.md",
        f"""# Feature and route inventory

**Last reviewed:** {TODAY}

## Primary navigation

| Nav label | Destination |
|---|---|
| Overview | `/dashboard/` |
| Transactions | `/budgets/transactions/` |
| Savings | `/savings/` |
| Pools | `/savings/pools/` |
| Investments | `/investing/` |
| Financial Goals | `/savings/goals/` |
| Insights | `/insights/` |
| Settings | `/accounts/settings/` |

## Route map (verified)

### Core
- `GET /` → `landing`
- `GET /dashboard/` → `dashboard` (login)
- `GET /insights/` → `insights` (login)

### Accounts
- `/accounts/signup|login|logout/`
- `/accounts/verify/`, `verify/resend/`, `verify/refresh/`
- `/accounts/onboarding/`, `onboarding/preferences/`
- `/accounts/manual-setup/`, `manual-setup/summary/`
- `/accounts/settings/`, `settings/messaging/link|unlink/`
- `/accounts/auth0/login/`, `/accounts/auth0/callback/`, root `/callback/`
- `GET /api/me/` (Bearer Auth0 access token)

### Budgets / transactions
- `/budgets/` index; periods/limits/income/expenses CRUD
- `/budgets/transactions/` list, add, edit, delete, CSV export
- `/budgets/categories/` add/delete

### Statements
- `/statements/`, `/statements/<pk>/`

### Imports
- `/imports/` history; `upload/`; `<pk>/account|preview|retry|delete/`

### Savings
- `/savings/`, `plan/` (+ refresh/status/allocations), `debts/` (+ add/edit)
- `pools/`, `goals/`, `add/`, `<pk>/`, edit, delete

### Investing
- `/investing/` portfolio tools
- demo budget add; demo-ai generate/accept/decline
- sweep; surplus propose; order approve/reject
- `market-data/`, `historical-demo/`, live-quotes JSON endpoints

### Assistant
- `POST /assistant/photon-webhook/`

### Admin
- `/admin/`

## Feature → code map

| Feature | Key files |
|---|---|
| Dashboard | `core/views.py`, `templates/core/dashboard.html` |
| Budgets | `budgets/services.py`, `budgets/views.py` |
| Transactions | `transactions/services.py`, `budgets/views.py` (UI) |
| Imports | `imports/pipeline.py`, `imports/views.py` |
| Savings plan | `savings/plan_engine.py`, `docs/savings-plan-engine.md` |
| Goals | `savings/services.py`, `savings/views.py` |
| Investing | `investing/services.py`, `investing/live_quotes.py` |
| Auth | `accounts/auth0.py`, `supabase_auth.py`, `identity.py` |
| Assistant | `assistant/services.py`, `intents.py`, `spectrum.py` |

## Incomplete / stub features

| Item | Evidence |
|---|---|
| Bank connect | Onboarding “coming soon”; `BANK_PROVIDER_API_KEY` unused |
| Real brokerage | Investing fully simulated; `BROKERAGE_PROVIDER_API_KEY` unused |
| Notification delivery | `NotificationEvent.delivery_status` stays `recorded` |
| Outbound iMessage | Webhook `delivered: false`; Spectrum Bun app is echo |
| Non-CSV imports | Only `csv` in `SOURCE_FORMAT_CHOICES` |
| Pagination | Not on transaction/ledger lists |
| Budget period rollover | Listed as next step in root README |
""",
    )

    # User guide
    write(
        DOCS / "user-guide/README.md",
        f"""# User guide

**Last reviewed:** {TODAY}

This guide explains how to operate Finovate as an end user. Screenshots use synthetic demo data (`demo@example.com`).

## Sections

1. [Landing and navigation](landing-and-navigation.md)
2. [Dashboard overview](dashboard.md)
3. [Transactions and statements](transactions.md)
4. [CSV import](imports.md)
5. [Budgets](budgets.md)
6. [Savings, pools, goals, and plan](savings.md)
7. [Simulated investing](investing.md)
8. [Insights](insights.md)
9. [Settings](settings.md)

For registration and login, see [Authentication](../authentication/README.md).
""",
    )

    write(
        DOCS / "user-guide/landing-and-navigation.md",
        f"""# Landing and navigation

{md_img("assets/screenshots/annotated/landing-annotated.png", "Annotated landing page")}

## Where to find it

- Public URL: `/`
- Authenticated users are directed into the app (dashboard / post-login destination)

## Interface elements

| Element | Behavior |
|---|---|
| Brand / logo | Product identity (hero-level on landing) |
| Sign in | Navigates to `/accounts/login/` |
| Sign up | Navigates to `/accounts/signup/` |
| Primary CTA | Starts signup / product entry |

## In-app navigation

After login, the left sidebar (desktop) or mobile menu lists Overview, Transactions, Savings, Pools, Investments, Financial Goals, Insights, Settings. Footer status shows “No bank connected” when no live bank integration is present.

{md_img("assets/screenshots/mobile/mobile-navigation.png", "Mobile navigation")}

## Responsive behavior

- Desktop (~1440px): persistent sidebar
- Mobile (~390px): compact header + toggle menu

## Errors / edge cases

- Without Auth0/Supabase configuration, login/signup pages render but cannot complete identity creation (messages instruct configuration / email sign-in availability).
""",
    )

    write(
        DOCS / "user-guide/dashboard.md",
        f"""# Dashboard (Overview)

{md_img("assets/screenshots/annotated/dashboard-overview-annotated.png", "Annotated dashboard")}

## Purpose

Home view of recorded cash, monthly income/spending, potential savings, recent transactions, budget progress snippets, and a savings-opportunity panel.

## Location

`/dashboard/` — requires login; typically after onboarding.

## Controls

| Control | Action |
|---|---|
| Chart range 1W/1M/3M/6M/1Y | Changes cash chart window |
| View all (transactions) | Goes to transaction list |
| Review savings plan | Opens savings plan / related CTA |
| Notification bell | UI affordance (events recorded in settings) |
| Log out | Ends session (+ Auth0 logout when applicable) |

## Data shown

- **Total cash** — income minus expenses (all time, recorded)
- **Monthly income / spending** — current month aggregates
- **Potential savings / investable surplus** — estimate after reserves, goals, bills buffer
- **Simulated portfolio** — kept separate from cash figure (disclaimer in UI)

## Failure / empty states

- Empty accounts show zeroed metrics and prompts to add data or import
- Manual-estimate banner may appear when onboarding used manual setup

## Acceptance checks

1. With seeded demo data, cash cards render non-zero values
2. Chart range buttons change the visible series without full logout
3. Footer disclaimer remains visible
""",
    )

    write(
        DOCS / "user-guide/transactions.md",
        f"""# Transactions and statements

{md_img("assets/screenshots/annotated/transactions-list-annotated.png", "Annotated transactions list")}

## Transactions list

- Path: `/budgets/transactions/`
- Search/filter, add/edit/delete, CSV export (`/budgets/transactions/export/`)

### Workflow: add a transaction

1. Start at Transactions
2. Open **Add**
3. Provide kind (income/expense), amount (positive), description, date, optional category
4. Submit → row appears; dashboard/budget progress recalculate from live sums
5. On validation failure, form redisplays with errors (no row created)

### Business rules

- Amounts are always stored positive; `kind` determines sign for display (`signed_amount`)
- Money math uses `Decimal` with `ROUND_HALF_UP` to 2 places (`core.services.quantize_money`)
- Per-user isolation: every query filters `user=request.user`

## Statements

{md_img("assets/screenshots/raw/statement-detail.png", "Statement detail")}

- `/statements/` lists statement accounts (e.g. Peachtree demo)
- Detail shows running balances and reconciliation vs stated totals
- Seed: `python manage.py seed_peachtree_demo`

## Edge cases

- Duplicate imported `external_id` does not double-count (`transactions.services.record_transaction`)
- No pagination yet — large histories may load slowly
""",
    )

    write(
        DOCS / "user-guide/imports.md",
        f"""# CSV bank-statement import

{md_img("assets/screenshots/annotated/imports-upload-annotated.png", "Import upload")}

## Purpose

Import real CSV statement rows into ordinary `Transaction` records after review. No fabricated amounts.

## Wizard

1. `/imports/upload/` — choose CSV (≤ 2MB)
2. `/imports/<id>/account/` — choose or create statement account
3. `/imports/<id>/preview/` — review detected columns/rows; select rows
4. Commit → status `committed`; transactions created with import source
5. History at `/imports/`; retry/delete available

## Technical notes

- Pipeline: `imports/pipeline.py` (+ enhanced CSV helpers when present)
- Model: `ImportBatch` stores raw file in `BinaryField`, column mapping, counters
- Unreadable files route toward manual setup guidance
- Only CSV is supported today

## Errors

| Condition | Behavior |
|---|---|
| File too large / invalid | Validation error; batch may be failed/discarded |
| Duplicate fingerprints | Skipped / not double-imported |
| Commit failure | Error message on batch; retry path |

See also existing deep dive: [bank-statement-import.md](../bank-statement-import.md).
""",
    )

    write(
        DOCS / "user-guide/budgets.md",
        f"""# Budgets

{md_img("assets/screenshots/raw/budgets-index.png", "Budgets index")}

## Purpose

Plan monthly income and recurring expenses, set category limits for a budget period, and see live progress.

## Location

`/budgets/` (also surfaced via Insights nav mapping in places)

## Workflows

### Create period and limits

1. Add a budget period (start/end, active flag)
2. Add category limits for that period
3. Progress computes **live** from expense transactions in-range (never stored spent totals)

### Income and recurring expenses

- Add/delete income sources and recurring expenses with cadence
- Cadence → monthly factor: weekly 4.33, biweekly 2.17, monthly 1, irregular 1

### Formulas (verified)

```
spent = sum(expenses in period for category)
remaining = limit - spent
percent_used = floor((spent / limit) * 100)   # int conversion as implemented
is_warning = percent_used >= warning_threshold and percent_used < 100
is_over = percent_used >= 100
```

Budget plan (`build_budget_plan`):

```
monthly_income = Σ monthly_equivalent(income)
monthly_recurring = Σ monthly_equivalent(expenses)
suggested_monthly_savings = heuristic from remaining goals / 12 (see services)
leftover_after_plan = income - recurring - suggested_savings
```

## Permissions

Login required; all objects scoped to the user.

## Acceptance

- Changing transactions updates progress without editing the Budget row
- Warning/over flags match threshold
""",
    )

    write(
        DOCS / "user-guide/savings.md",
        f"""# Savings, pools, goals, and plan

{md_img("assets/screenshots/raw/savings-index.png", "Savings index")}

## Goals and pools

- **Financial Goals** `/savings/goals/` — deadline-based goals with pace suggestions
- **Pools** `/savings/pools/` — investable / system pools
- Goal detail supports contributions/withdrawals/adjustments

Milestones at 25/50/75/100% create `NotificationEvent` records (not externally delivered).

{md_img("assets/screenshots/raw/savings-goals.png", "Goals list")}

## Personalized savings plan

{md_img("assets/screenshots/annotated/savings-plan-annotated.png", "Savings plan annotated")}

Recommendation-only engine (`savings/plan_engine.py`):

```
Monthly Surplus = Income − Expenses
Safe-to-Save = max(0, Surplus − additional buffer)
```

Scenarios auto-selected: stabilization, build_reserve, debt_focus, stable, strong.

Default system pools (created if missing): Rainy Day, Investment, Regular Savings.

UI:

- `/savings/plan/` full explanation, warnings, history
- Refresh / pause / disable / edit allocation percents (must total 100% of allocated safe-to-save)
- `/savings/debts/` optional APR/balance inputs for debt_focus

Never transfers money or places trades.

Details: [savings-plan-engine.md](../savings-plan-engine.md).
""",
    )

    write(
        DOCS / "user-guide/investing.md",
        f"""# Simulated investing

{md_img("assets/screenshots/annotated/investing-index-annotated.png", "Investing annotated")}

## Critical disclaimer

**No brokerage is connected. No real orders are placed.** UI labels this throughout.

## Features

| Feature | Path / action | Behavior |
|---|---|---|
| Portfolio tab | `/investing/` | Snapshots, pending round-ups, surplus propose |
| Round-ups | settings + sweep | Expense → pending ledger → simulated sweep |
| Demo budget | add virtual funds | Caps demo AI / explore buys |
| Demo AI ideas | generate/accept/decline | Rule-based; prefers Twelve Data history |
| Explore | `/investing/market-data/` | Sample/live quotes; simulated buys |
| Historical demo | `/investing/historical-demo/` | Hypothetical scenarios on verified prices |
| Order approve/reject | POST endpoints | Fills or rejects `SimulatedOrder` |

## Round-up math (verified)

```
round_up = round_up_to_nearest(expense_amount, round_to)
# Example: 4.35 with round_to 1.00 → 0.65
```

Sweep respects pause flag, contribution cap, and minimum cash reserve.

## Live quotes

- Polling JS: `static/js/live-holdings.js`, `live-market-quotes.js` (~5 min)
- Requires `TWELVE_DATA_API_KEY`; otherwise UI explains history/quotes unavailable

## Errors

- Propose without surplus / over budget → validation/message
- Twelve Data credit exhaustion → usage window blocks; cached quotes may remain
""",
    )

    write(
        DOCS / "user-guide/insights.md",
        f"""# Insights

{md_img("assets/screenshots/raw/insights-overview.png", "Insights page")}

## Location

`/insights/`

## Contents

- Month-over-month spending bars
- Category donut / breakdown
- Savings rate style metrics
- Recurring bills summary

Derived from the user’s transactions and budget-related records. Login required.
""",
    )

    write(
        DOCS / "user-guide/settings.md",
        f"""# Settings

{md_img("assets/screenshots/annotated/settings-profile-annotated.png", "Settings annotated")}

{md_img("assets/screenshots/raw/theme-dark-mode.png", "Dark theme preview")}

## Location

`/accounts/settings/`

## Controls

| Area | Fields / actions |
|---|---|
| Preferences | Income cadence, risk comfort, cash reserve target, budget warning % |
| Theme | `light` / `dark` / `system` → `data-theme` on html/body |
| Notifications | Toggles for budget warnings, goal milestones, simulated orders |
| Messaging | Issue / unlink Photon link code |
| Connected accounts | Placeholders (bank/brokerage) — status only |
| Notification events | In-app recorded list |

## Persistence

Saved on `accounts.Profile`. Theme applies immediately via template context.
""",
    )

    # Features index + individual specs
    write(
        DOCS / "features/README.md",
        f"""# Feature specifications

**Last reviewed:** {TODAY}

| Feature | Spec |
|---|---|
| Dashboard | [dashboard.md](dashboard.md) |
| Transactions | [transactions.md](transactions.md) |
| CSV import | [csv-import.md](csv-import.md) |
| Budgets | [budgets.md](budgets.md) |
| Savings goals & pools | [savings-goals.md](savings-goals.md) |
| Savings plan engine | [savings-plan.md](savings-plan.md) |
| Simulated investing | [investing.md](investing.md) |
| Market data explorer | [market-data.md](market-data.md) |
| Insights | [insights.md](insights.md) |
| Messaging assistant | [messaging-assistant.md](messaging-assistant.md) |
| Manual setup | [manual-setup.md](manual-setup.md) |
""",
    )

    def feature_page(name, path, purpose, users, deps, ui, workflow, logic, impl, errors, tests, shot=None):
        shot_block = md_img(shot, name) if shot else ""
        write(
            DOCS / "features" / path,
            f"""# {name}

**Last reviewed:** {TODAY}

## 1. Feature overview

- **Purpose:** {purpose}
- **Intended users:** {users}
- **Permissions:** Authenticated user owning the data (unless noted)
- **Dependencies / prerequisites:** {deps}
- **Related:** See [feature inventory](../feature-inventory.md)

{shot_block}

## 2. User interface walkthrough

{ui}

## 3. Step-by-step workflow

{workflow}

## 4. Business logic

{logic}

## 5. Technical implementation

{impl}

## 6. Errors and edge cases

{errors}

## 7. Testing and acceptance criteria

{tests}
""",
        )

    feature_page(
        "Dashboard",
        "dashboard.md",
        "Summarize recorded finances and surface next actions.",
        "Signed-in users who finished (or skipped into) the app",
        "Transactions/budgets/goals data improve usefulness; works empty",
        "Overview cards, cash chart ranges, recent txns, savings opportunity, insights snippets. "
        + md_img("assets/screenshots/annotated/dashboard-overview-annotated.png", "Dashboard"),
        "1. Land on `/dashboard/` after login.\n2. Review cards and chart.\n3. Jump to transactions, savings plan, or investing as needed.\n4. If operation fails (session expired) → redirect to login.",
        "Aggregations over user’s transactions; surplus estimate subtracts bills buffer, goal allocations, reserve. Simulated portfolio kept separate.",
        "Frontend: `templates/core/dashboard.html`, `static/js/finovate.js`. Backend: `core/views.py`. Data: Transaction, Budget, Goal, PortfolioSnapshot, Profile.",
        "Expired session → login redirect. Empty data → zeros/empty lists. No bank connected badge always when integration absent.",
        "Seed demo → open dashboard → non-zero cash; chart controls respond; disclaimer visible.",
        "assets/screenshots/raw/dashboard-overview.png",
    )

    feature_page(
        "Transactions",
        "transactions.md",
        "Record and manage income/expense rows powering budgets and insights.",
        "Signed-in users",
        "Optional categories; optional statement accounts",
        "List with search; add/edit forms; delete confirm; CSV export. "
        + md_img("assets/screenshots/raw/transactions-list.png", "Transactions"),
        "Add → validate → save → redirect list. Edit/delete similar. Export downloads CSV of user’s rows.",
        "Positive `amount` + `kind`; `quantize_money`; idempotent import via `external_id`.",
        "`budgets/views.py` (UI), `transactions/models.py`, `transactions/services.py`, forms in budgets.",
        "Invalid amount/date → form errors. Cross-user pk → 404. Duplicate external_id → no double count.",
        "CRUD isolation tests in `budgets/tests.py` / `transactions/tests.py`.",
    )

    feature_page(
        "CSV import",
        "csv-import.md",
        "Bring bank CSV history into Finovate without inventing rows.",
        "Signed-in users during or after onboarding",
        "CSV ≤ 2MB; readable columns",
        "Upload → account → preview → commit. " + md_img("assets/screenshots/raw/imports-upload.png", "Upload"),
        "See user-guide/imports.md wizard steps; cancel/delete batch; retry failed.",
        "Column detection, dedupe fingerprints, commit creates Transactions with source imported.",
        "`imports/views.py`, `pipeline.py`, `models.ImportBatch`.",
        "Unreadable → guidance/manual setup. Oversized file rejected. Commit errors stored on batch.",
        "`imports/tests.py` pipeline + isolation coverage.",
    )

    feature_page(
        "Budgets",
        "budgets.md",
        "Set limits and see live spending progress.",
        "Signed-in users",
        "Categories + active BudgetPeriod",
        md_img("assets/screenshots/raw/budgets-index.png", "Budgets"),
        "Create period → add limits → spend via transactions → progress updates. Delete limit/period via confirm.",
        "See user-guide/budgets.md formulas; warning thresholds per budget or profile default.",
        "`budgets/services.py`, models BudgetPeriod/Budget/IncomeSource/RecurringExpense.",
        "Overlapping unique (period, category). Zero limit → percent_used 0.",
        "Progress/plan tests in `budgets/tests.py`.",
    )

    feature_page(
        "Savings goals and pools",
        "savings-goals.md",
        "Track targets, contributions, and investable pools.",
        "Signed-in users",
        "None beyond login",
        md_img("assets/screenshots/raw/savings-goals.png", "Goals"),
        "Add goal → contribute → milestones fire events → edit/delete. Pools page filters investable/system goals.",
        "`progress_percent`, `remaining_amount`; pace estimates; withdrawal cannot go negative.",
        "`savings/services.py`, `Goal`, `Contribution`.",
        "Over-withdraw blocked. Overdue goals prompt new date.",
        "`savings/tests.py`.",
    )

    feature_page(
        "Savings plan engine",
        "savings-plan.md",
        "Recommend how to allocate safe-to-save across pools/debt.",
        "Signed-in users with income/expense context",
        "Income sources, expenses, optional debts, emergency balances",
        md_img("assets/screenshots/raw/savings-plan.png", "Plan"),
        "Open `/savings/plan/` → review scenario → optional refresh/status/allocations → edit debts as needed.",
        "Scenario table and safe-to-save math in `docs/savings-plan-engine.md` (verified against `plan_engine.py`).",
        "`savings/plan_engine.py`, `SavingsPlan`, `SavingsPlanRevision`, `DebtObligation`.",
        "No income → stabilization. Unrealistic goals flagged. Fingerprint skips noisy rewrites.",
        "`savings/test_plan_engine.py`.",
    )

    feature_page(
        "Simulated investing",
        "investing.md",
        "Educate with round-ups, surplus proposals, and demo orders — without real trades.",
        "Signed-in users",
        "Optional Twelve Data key for live quotes",
        md_img("assets/screenshots/annotated/investing-index-annotated.png", "Investing"),
        "Enable round-ups → spend → sweep → propose surplus → approve/reject orders → explore market demo buys.",
        "`round_up_to_nearest`; sweep caps/reserves; demo AI ranking rules; portfolio snapshots.",
        "`investing/services.py`, `live_quotes.py`, models RoundUp*, SimulatedOrder, DemoInvestment*.",
        "Paused round-ups skip sweep. Missing API key → sample/history messaging. Always simulated.",
        "`investing/tests.py` (large suite).",
    )

    feature_page(
        "Market data explorer",
        "market-data.md",
        "Browse sample/live prices and place budget-limited simulated buys.",
        "Signed-in users",
        "Seeded MarketPrice rows; optional Twelve Data",
        md_img("assets/screenshots/raw/investing-market-data.png", "Market data"),
        "Open Explore → inspect chart → invest amount within demo budget → order simulated.",
        "Uses verified OHLCV sample and/or Twelve Data; charts via `market-charts.js`.",
        "`investing/views.py`, `MarketPrice`, `LiveMarketQuote`.",
        "Insufficient demo budget → rejected. Stale quotes possible when credits exhausted.",
        "Market explorer tests in `investing/tests.py`.",
    )

    feature_page(
        "Insights",
        "insights.md",
        "Visualize spending patterns and savings rate style metrics.",
        "Signed-in users",
        "Transaction history",
        md_img("assets/screenshots/raw/insights-overview.png", "Insights"),
        "Open `/insights/` → review MoM and category visuals.",
        "Aggregations in `core/views.py` / display helpers.",
        "`templates/core/insights.html`, `core/display.py`.",
        "Empty history → empty/zero charts.",
        "Covered partly by `core/tests.py` empty states.",
    )

    feature_page(
        "Messaging assistant",
        "messaging-assistant.md",
        "Answer finance questions / gated actions via Spectrum webhook.",
        "Users who linked messaging in Settings",
        "Webhook secret; link code; Spectrum project",
        "No primary web page — Settings messaging link + webhook. See [messaging-assistant.md](../messaging-assistant.md).",
        "Issue link code → user texts code → chat binds `photon_chat_id` → intents handled → confirmations for sensitive actions.",
        "Keyword/regex intents in `assistant/intents.py`; reject withdraw invention etc.",
        "`assistant/views.py`, `services.py`, `spectrum.py`, models PendingConfirmation, MessagingLinkCode.",
        "Bad signature → reject. Idempotent delivery keys. Django replies `delivered: false`.",
        "`assistant/tests.py`; local `manage.py chat_demo`.",
    )

    feature_page(
        "Manual setup",
        "manual-setup.md",
        "Bootstrap estimates when CSV/bank unavailable.",
        "Users in onboarding",
        "Onboarding path",
        md_img("assets/screenshots/raw/manual-setup.png", "Manual setup"),
        "Enter savings balance / monthly goal style fields → summary via `build_budget_plan` → continue preferences.",
        "`accounts/manual_setup.py` applies profile fields; does not invent transactions.",
        "`accounts/views.py` manual_setup*, Profile manual_* fields.",
        "Invalid numbers → form errors.",
        "Accounts tests for manual setup.",
    )

    # Workflows
    write(
        DOCS / "workflows/README.md",
        f"""# Workflows

**Last reviewed:** {TODAY}

| Workflow | Doc |
|---|---|
| Registration & verification | [registration.md](registration.md) |
| Login & session | [login-session.md](login-session.md) |
| Onboarding | [onboarding.md](onboarding.md) |
| CSV import | [csv-import.md](csv-import.md) |
| Round-up sweep | [roundup-sweep.md](roundup-sweep.md) |
| Savings plan refresh | [savings-plan-refresh.md](savings-plan-refresh.md) |
""",
    )

    write(
        DOCS / "workflows/registration.md",
        f"""# Registration and email verification

```mermaid
sequenceDiagram
  participant U as User
  participant F as Finovate
  participant A as Auth0/Supabase
  U->>F: GET /accounts/signup/
  alt Auth0 configured
    U->>F: POST name/email/phone
    F->>A: Universal Login signup
    A->>F: /callback with code
    F->>F: Profile by auth0_sub
  else Supabase configured
    U->>F: POST email/password
    F->>A: sign_up
    A->>F: user / confirm email
    F->>F: Profile by supabase_uid
  end
  alt email unverified (Auth0)
    F->>U: /accounts/verify/
  else onboarding incomplete
    F->>U: /accounts/onboarding/
  else
    F->>U: /dashboard/
  end
```

### Failure paths

- Rate limit signup 10 / 10 min per IP
- Auth0 API errors → authentication failed messaging
- Weak Supabase password → form error
""",
    )

    write(
        DOCS / "workflows/login-session.md",
        f"""# Login and session management

1. User opens `/accounts/login/`
2. Auth0 button/flow or Supabase email POST
3. On success: Django session login (`ModelBackend`); Auth0 sets session flags + cookies via SDK
4. `Auth0SessionMiddleware` may force verify page for unverified Auth0 users
5. `destination_for` routes next
6. Logout clears Django session; Auth0 federated logout when `logged_in_via_auth0`

Rate limit: 20 login attempts / 10 minutes / IP.
""",
    )

    write(
        DOCS / "workflows/onboarding.md",
        f"""# Onboarding

{md_img("assets/screenshots/raw/onboarding-start.png", "Onboarding")}

1. `/accounts/onboarding/` — choose Import CSV / Bank (coming soon) / Continue
2. Optional import wizard or `/accounts/manual-setup/` → summary
3. `/accounts/onboarding/preferences/` — cadence, risk, reserve, notifications, phone/consent
4. Sets `Profile.onboarding_completed = True`
5. Redirect dashboard
""",
    )

    write(
        DOCS / "workflows/csv-import.md",
        """# CSV import workflow

1. Start: `/imports/upload/` or onboarding import
2. Upload CSV → `ImportBatch` draft with raw bytes
3. Choose/create `Account`
4. Preview mapped rows; deselect bad rows
5. Commit → transactions + counters; `committed_at`
6. Alternate: retry, delete (may keep committed txns while dropping file)
""",
    )

    write(
        DOCS / "workflows/roundup-sweep.md",
        """# Round-up sweep workflow

1. Enable round-ups in Investing settings; set round_to, cap, reserve, target pool
2. Record expense transaction → pending `RoundUpLedgerEntry` (when rules match)
3. User clicks **Simulate a sweep now** or equivalent service path
4. System checks pause, reserve, cap
5. On success: entries marked swept; `PortfolioSnapshot` / simulated contribution updated
6. Failure: pause/cap/reserve → no sweep; user sees messaging
""",
    )

    write(
        DOCS / "workflows/savings-plan-refresh.md",
        """# Savings plan refresh

1. User opens `/savings/plan/` or triggers refresh
2. `analyze_finances` → `choose_scenario` → `build_recommendation`
3. If fingerprint/material change: write `SavingsPlan` + `SavingsPlanRevision`
4. Else: keep existing recommendation
5. User may pause/disable/override percents (100% constraint)
""",
    )

    # Auth / account / interface
    write(
        DOCS / "authentication/README.md",
        f"""# Authentication

**Last reviewed:** {TODAY}

{md_img("assets/screenshots/annotated/authentication-login-annotated.png", "Login")}

## Providers

| Provider | When | Password storage in Finovate |
|---|---|---|
| Auth0 | `AUTH0_DOMAIN/CLIENT_ID/CLIENT_SECRET/SECRET` set | None — profile keyed by `auth0_sub` |
| Supabase Auth | `SUPABASE_URL` + publishable/anon key | None — keyed by `supabase_uid` |

## Pages

- Signup `/accounts/signup/`
- Login `/accounts/login/`
- Verify `/accounts/verify/` (+ resend/refresh)
- Logout `/accounts/logout/`
- Callback `/callback/`

## Identity rules (verified)

- Auth0 profile lookup **by subject only**, never email merge
- Unverified Auth0 users gated by middleware
- Optional Management API for verification email resend (`AUTH0_MGMT_*`)
- Legacy Django-password users can migrate to Supabase on first email login when service role configured

Deep dive: [architecture-auth0-supabase.md](../architecture-auth0-supabase.md), [auth0-verification-supabase-backup.md](../auth0-verification-supabase-backup.md).
""",
    )

    write(
        DOCS / "account-management/README.md",
        f"""# Account management

**Last reviewed:** {TODAY}

## Profile

Auto-created on User create (`accounts/signals.py`). Stores preferences, theme, Auth0/Supabase ids, messaging fields, manual setup estimates, onboarding flag.

## Settings lifecycle

See [user-guide/settings.md](../user-guide/settings.md).

## Notification events

Types include budget_warning, goal_milestone, roundup_paused, account_created, welcome, system. **Recorded only** — no email/push/SMS sender.

## Connected accounts

`ConnectedAccountPlaceholder` — metadata statuses only; no credentials stored.
""",
    )

    write(
        DOCS / "interface/README.md",
        f"""# Interface

**Last reviewed:** {TODAY}

## Stack

Server-rendered Django templates + `/static/css/finovate.css` (DM Sans / Manrope via Google Fonts import) + progressive JS (`finovate.js`, live quote pollers, market charts).

## Theme

Profile `theme`: light / dark / system → `data-theme` attribute. Dark capture:

{md_img("assets/screenshots/raw/theme-dark-mode.png", "Dark theme")}

## Navigation

Sidebar desktop; mobile toggle. Brand logo partial `partials/brand_logo.html`.

## Motions / interactions

- Mobile menu open/close
- Review modal (client-side; does not persist a plan)
- Surplus allocation slider / portfolio projection calculator in `finovate.js`
- Chart scrubber on Explore

## Responsive

Desktop and mobile screenshots under `assets/screenshots/`.

{md_img("assets/screenshots/mobile/mobile-dashboard.png", "Mobile dashboard")}
""",
    )

    write(
        DOCS / "api/README.md",
        f"""# API reference

**Last reviewed:** {TODAY}

Finovate is primarily form-posted HTML. JSON endpoints:

## `GET /api/me/`

- **Auth:** `Authorization: Bearer <Auth0 access token>` with audience `AUTH0_AUDIENCE`
- **Success:** `{{ "identity_status", "email_verified", "transaction_count" }}`
- **Errors:** 401 invalid/missing token; 403 when appropriate

## `POST /assistant/photon-webhook/`

- **Auth:** Spectrum HMAC signature or legacy webhook secret
- **Body:** Photon webhook payload
- **Response:** JSON including `reply`, `delivered: false` (Django does not send iMessage)
- **Idempotency:** `SpectrumDelivery` unique `(webhook_id, message_id)`

## `GET /investing/live-quotes/`

- **Auth:** Django session login
- **Response:** JSON quotes for user’s filled demo holdings

## `GET /investing/market-data/live-quotes/`

- **Auth:** Django session login
- **Response:** JSON shared market stock quotes

## Form POST endpoints

All feature mutations (transactions, budgets, goals, investing actions, imports) are session-authenticated form posts with CSRF, returning redirects/HTML — not a public REST surface.
""",
    )

    write(
        DOCS / "database/README.md",
        f"""# Database reference

**Last reviewed:** {TODAY}

Default: SQLite (`db.sqlite3` local; `/tmp/finovate.sqlite3` on Vercel). Optional Postgres via `DATABASE_URL` / `SUPABASE_DB_URL`.

## ER overview

```mermaid
erDiagram
  User ||--|| Profile : has
  User ||--o{{ Transaction : owns
  User ||--o{{ Category : owns
  User ||--o{{ Account : owns
  User ||--o{{ BudgetPeriod : owns
  BudgetPeriod ||--o{{ Budget : contains
  Budget }}o--|| Category : limits
  User ||--o{{ Goal : owns
  Goal ||--o{{ Contribution : has
  User ||--o| SavingsPlan : has
  SavingsPlan ||--o{{ SavingsPlanRevision : history
  User ||--o{{ DebtObligation : owns
  User ||--o{{ ImportBatch : owns
  User ||--o| RoundUpSettings : has
  User ||--o{{ RoundUpLedgerEntry : has
  User ||--o{{ SimulatedOrder : has
  Transaction ||--o| RoundUpLedgerEntry : optional
```

## Core tables (fields summarized)

### Profile (`accounts_profile`)
Preferences, theme, notification toggles, `auth0_sub`, `supabase_uid`, identity/verification flags, messaging fields, manual setup amounts, `onboarding_completed`.

### NotificationEvent / ConnectedAccountPlaceholder
Event log; integration placeholders.

### Account / Category / Transaction (`transactions_*`)
Statement metadata; categories income/expense; transactions with optional account/category, `external_id` uniqueness, `source` manual|imported|demo.

### Budgets
IncomeSource, RecurringExpense, BudgetPeriod, Budget (unique period+category).

### Savings
Goal (incl. system_key / investable flags), Contribution, DebtObligation, SavingsPlan, SavingsPlanRevision.

### Investing
DemoInvestmentFunding, LiveMarketQuote, TwelveDataUsageWindow, TwelveDataHistoryState, RoundUpSettings, RoundUpLedgerEntry, PortfolioSnapshot, SimulatedOrder, DemoInvestmentSuggestion, MarketPrice.

### Imports
ImportBatch — raw file, status draft|committed|failed|discarded, mapping, counters.

### Assistant
PendingConfirmation, SpectrumDelivery, MessagingLinkCode, ConversationMessage.

## Migrations

Per-app under `*/migrations/`. Apply with `python manage.py migrate`.

## Retention / deletion

- User-owned rows cascade with user/FKs as modeled
- Import delete may remove file bytes while keeping committed transactions (see imports views)
- No automated TTL job in-repo
""",
    )

    write(
        DOCS / "architecture/README.md",
        f"""# Architecture

**Last reviewed:** {TODAY}

## Request flow

```mermaid
flowchart TD
  Req[HTTP request] --> MW[Django middleware + Auth0SessionMiddleware]
  MW --> URL[URLConf]
  URL --> View[Thin view]
  View --> Svc[services.py / plan_engine / pipeline]
  Svc --> ORM[Django ORM]
  ORM --> DB[(SQLite/Postgres)]
  View --> Tpl[Template + context processors]
  Tpl --> Res[HTML response]
```

## Design invariants

1. **Decimal money only** — `core.services`
2. **Business logic in services** — reusable by web + assistant
3. **Per-user ownership** on every query
4. **Simulated vs real labeled** in investing UI
5. **IdP owns passwords** — Finovate stores subjects/uids

## Deployment architecture

```mermaid
flowchart LR
  User --> Vercel
  Vercel --> WSGI[finovate/wsgi.py]
  WSGI --> Migrate[migrate on boot]
  WSGI --> App[Django]
  App --> PG[(Postgres optional)]
  App --> Auth0
  App --> Twelve[Twelve Data]
  Spectrum[Spectrum sidecar] -. webhook .-> App
```

On Vercel: `VERCEL` enables hosts/SSL/tmp SQLite unless Postgres URL set; optional `seed_demo` when `DEMO_ACCOUNT_PASSWORD` set and no demo user.
""",
    )

    write(
        DOCS / "security/README.md",
        f"""# Security

**Last reviewed:** {TODAY}

## Controls (verified)

| Control | Implementation |
|---|---|
| AuthN | Auth0 OIDC / Supabase Auth |
| AuthZ | Login required + user-scoped ORM filters |
| CSRF | Django middleware on form posts |
| Sessions | Django sessions; Auth0 SDK cookies when used |
| Rate limits | Cache IP limits on signup/login/resend |
| Webhook auth | Spectrum HMAC / shared secret |
| API tokens | Auth0 JWT via JWKS (`accounts/tokens.py`) |
| Secrets | Env vars only; `.env.example` empty placeholders |
| Money integrity | Decimal quantization |

## Validation

Django forms for user input; import pipeline validates CSV structure/size.

## Known gaps / hardening

- CDN fonts without SRI (noted in README)
- No CSS/JS build pipeline / minification
- Notification channel not secured because not delivered
- Bank/brokerage keys are placeholders — must not be treated as live integrations
""",
    )

    write(
        DOCS / "integrations/README.md",
        f"""# Integrations

**Last reviewed:** {TODAY}

| Integration | Purpose | Config | Failure behavior |
|---|---|---|---|
| Auth0 | Primary identity | `AUTH0_*`, `APP_BASE_URL` | Signup/login unavailable if incomplete |
| Supabase Auth | Email/password | `SUPABASE_URL`, keys | Forms hidden if unconfigured |
| Supabase/Postgres | Hosted DB | `DATABASE_URL` / `SUPABASE_DB_URL` | Falls back to SQLite |
| Twelve Data | Live quotes/history | `TWELVE_DATA_API_KEY` | Sample data / unavailable messaging; credit windows |
| Photon Spectrum | Inbound messaging | `IMESSAGE_*`, webhook secrets | Invalid signature rejected; no Django outbound |
| Bank provider | Future | `BANK_PROVIDER_API_KEY` | Unused |
| Brokerage | Future | `BROKERAGE_PROVIDER_API_KEY` | Unused |
| Email | Django console backend | — | Prints to console in dev |

See also: [integration-readiness.md](../integration-readiness.md), [messaging-assistant.md](../messaging-assistant.md).
""",
    )

    write(
        DOCS / "troubleshooting.md",
        f"""# Troubleshooting

**Last reviewed:** {TODAY}

| Symptom | Likely cause | Resolution |
|---|---|---|
| Login/signup does nothing useful | Auth0/Supabase not configured | Set env vars; restart server |
| Stuck on verify page | Auth0 email unverified | Resend email; confirm in Auth0; refresh |
| Callback error | `APP_BASE_URL` mismatch | Align Auth0 callback with `<APP_BASE_URL>/callback` |
| Empty dashboard after seed | Viewing wrong user / not logged in | Use configured IdP account; seed is fixture |
| Import rejected | >2MB or unreadable CSV | Shrink/fix file; use manual setup |
| Twelve Data warnings | Missing/exhausted API key | Set key or rely on sample market data |
| Investing “orders” not in brokerage | Expected | Simulated only |
| Webhook 401/403 | Bad signature/secret | Check `IMESSAGE_WEBHOOK_SECRET` / Photon secret |
| Vercel data loss | Ephemeral `/tmp` SQLite | Configure Postgres URL |
| Tests fail on missing migrations | Models changed | `python manage.py makemigrations` / migrate in CI |
""",
    )

    write(
        DOCS / "development/README.md",
        f"""# Development

**Last reviewed:** {TODAY}

## Layout

See root README architecture table. Put domain logic in `services.py` (or `plan_engine.py` / `pipeline.py`), keep views thin.

## Commands

```bash
python manage.py runserver
python manage.py test
python manage.py test savings.investing  # example subset
python manage.py chat_demo
python manage.py seed_demo
```

## Contribution notes

- Never use floats for money
- Filter by `request.user` on every user data path
- Label simulated investing clearly
- Do not commit secrets
- Prefer extending tests beside features
""",
    )

    write(
        DOCS / "deployment/README.md",
        f"""# Deployment

**Last reviewed:** {TODAY}

## Vercel

- `vercel.json` sets Django framework + `finovate/wsgi.py` `maxDuration: 60`
- WSGI boot may migrate + seed demo when configured
- Prefer Postgres (`DATABASE_URL`) for persistence
- Set `FINOVATE_DEBUG=false`, strong `FINOVATE_SECRET_KEY`, Auth0, `APP_BASE_URL=https://your-domain`

## Checklist

1. Configure IdP callback URLs
2. Set CSRF trusted origins / allowed hosts (Vercel branch in settings helps `*.vercel.app`)
3. Configure webhook URL for Spectrum if used
4. Do not expose management/service-role keys to the browser
5. Confirm Twelve Data key only on server
""",
    )

    write(
        DOCS / "testing/README.md",
        f"""# Testing

**Last reviewed:** {TODAY}

## Strategy

Django `TestCase` / client tests per app. Run:

```bash
python manage.py test
```

Hundreds of `test_*` methods cover auth, imports, budgets, savings plan engine, investing, assistant, isolation.

## Regression priorities

1. Per-user isolation
2. Decimal round-up example $4.35 → $0.65
3. Idempotent imports
4. Budget progress live recompute
5. Savings plan scenarios + allocation caps
6. Webhook signature + confirmation gating
7. Simulated order approve/reject

## UI verification for docs

Playwright captures against local `runserver` with force-login session cookie for the `demo` fixture (documentation only — not a production login path).
""",
    )

    write(
        DOCS / "glossary.md",
        f"""# Glossary

**Last reviewed:** {TODAY}

| Term | Meaning |
|---|---|
| Safe-to-save | Surplus after expenses and buffers; plan allocation basis |
| Round-up | Cents added to reach next `round_to` multiple on an expense |
| Sweep | Moving pending round-ups into simulated portfolio/pool |
| Simulated order | DB-only investment order; not sent to a broker |
| Investable surplus | Cash minus reserve, goal allocations, bills buffer |
| System pool | Auto-created Rainy Day / Investment / Regular Savings goals |
| ImportBatch | One CSV upload lifecycle |
| auth0_sub | Auth0 subject identifier on Profile |
| Spectrum | Photon messaging adapter / webhook channel |
| Peachtree demo | Fictional statement seed for reconciliation demo |
""",
    )

    write(
        DOCS / "coverage-matrix.md",
        f"""# Documentation coverage matrix

**Last reviewed:** {TODAY}

| Feature | Doc page | Screenshot | Tests exist | Verified vs code | Status |
|---|---|---|---|---|---|
| Landing | user-guide/landing-and-navigation.md | raw+annotated+mobile | partial | Yes | Complete |
| Login/Signup UI | authentication/README.md | raw+annotated+mobile | Yes (accounts) | Yes | Complete (IdP config-dependent runtime) |
| Dashboard | features/dashboard.md | raw+annotated+mobile | Yes | Yes | Complete |
| Insights | features/insights.md | raw | partial | Yes | Complete |
| Transactions | features/transactions.md | raw+annotated+mobile | Yes | Yes | Complete |
| Statements | user-guide/transactions.md | raw | Yes | Yes | Complete |
| CSV import | features/csv-import.md | raw+annotated | Yes | Yes | Complete |
| Budgets | features/budgets.md | raw | Yes | Yes | Complete |
| Savings index/pools/goals | features/savings-goals.md | raw+mobile | Yes | Yes | Complete |
| Savings plan/debts | features/savings-plan.md | raw+annotated | Yes | Yes | Complete |
| Investing | features/investing.md | raw+annotated+mobile | Yes | Yes | Complete |
| Market data / historical | features/market-data.md | raw | Yes | Yes | Complete |
| Settings/theme | user-guide/settings.md | raw+annotated+dark+mobile | Yes | Yes | Complete |
| Onboarding/manual setup | features/manual-setup.md | raw | Yes | Yes | Complete |
| Messaging assistant | features/messaging-assistant.md | n/a (no page) | Yes | Yes | Complete (webhook; no UI shot) |
| API /api/me | api/README.md | n/a | Yes | Yes | Complete |
| Database/architecture | database/, architecture/ | diagrams | n/a | Yes | Complete |
| DOCX deliverable | deliverables/*.docx | embedded | n/a | Yes | Complete |
| PPT deliverable | deliverables/*.pptx | embedded | n/a | Yes | Complete |

### Screenshot checklist gaps (optional follow-up)

- Form validation error states (transaction add invalid)
- Import preview populated mid-wizard
- Auth0 Universal Login hosted page (third-party; not captured)
- Order approve confirmation flash messages
""",
    )

    write(
        DOCS / "audit-report.md",
        f"""# Documentation audit report

**Date:** {TODAY}
**Version:** {VERSION}

## Completed work

1. Full codebase route/model/feature inventory
2. `/docs` hierarchical documentation system
3. Per-feature specifications with UI, workflow, logic, implementation, errors, tests
4. Real screenshot capture (desktop, mobile, annotated) from local demo data
5. Architecture / ER / auth Mermaid diagrams
6. API, database, security, integrations, deployment, testing guides
7. Coverage matrix
8. Packaged **DOCX** and **PPTX** deliverables

## Features verified against source

Dashboard, transactions, statements, imports, budgets, savings goals/pools, savings plan engine, debts, investing (simulated), market data, insights, settings/themes, auth flows (code-level), assistant webhook, `/api/me/`, Vercel deployment hooks.

## Defects / discrepancies discovered (reported, not “fixed for docs”)

1. Root README still says “35 tests” and mentions Tailwind/Chart.js CDN stack; current UI is primarily `finovate.css` custom design and test count is much higher.
2. README identity wording has minor internal tension (Auth0-required vs local login mentions); implementation uses Auth0 and/or Supabase; `seed_demo` is not website login when IdPs are expected.
3. Bank/brokerage env keys unused (placeholders) — correctly labeled coming soon / simulated.
4. Notification delivery not implemented despite settings toggles.
5. Spectrum Bun sidecar is echo-only relative to Django assistant logic.
6. Review modal “save plan” is client-only hide/show (`finovate.js`), does not persist.
7. Transaction/ledger pagination absent.

## Known limitations of this documentation pass

- Screenshots used Django `force_login` session for the `demo` fixture because Auth0/Supabase were not configured in the capture environment.
- Auth0 hosted Universal Login UI not captured (external).
- Some mid-wizard import preview / validation-error screenshots listed as follow-up in coverage matrix.
- DOCX/PPT are comprehensive guides, not a live searchable web portal (Markdown tree is the navigable portal substitute without adding a heavy docs framework).

## Outstanding tasks

1. Capture validation-error and import-preview screenshots when convenient
2. Refresh root README stack/test-count wording in a separate non-docs PR if desired
3. Optional: MkDocs/Sphinx portal if product owners want search UI beyond Markdown + DOCX/PPT
""",
    )

    write(
        DOCS / "CHANGELOG.md",
        f"""# Documentation changelog

## {VERSION} — {TODAY}

- Initial comprehensive documentation system under `/docs`
- Feature inventory, user guides, workflows, API/DB/architecture/security
- Screenshot library: raw, annotated, mobile
- Deliverables: `Finovate-Complete-Documentation.docx`, `Finovate-Visual-Feature-Guide.pptx`
- Coverage matrix and audit report
""",
    )


def add_heading(doc: Document, text: str, level: int = 1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = NAVY
    return h


def add_para(doc: Document, text: str, bold: bool = False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(11)
    return p


def add_bullets(doc: Document, items: list[str]):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def try_picture(doc: Document, path: Path, width_in: float = 6.2):
    if path.exists():
        doc.add_picture(str(path), width=Inches(width_in))
        last = doc.paragraphs[-1]
        last.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return True
    return False


def generate_docx() -> Path:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_heading("Finovate", 0)
    for r in title.runs:
        r.font.color.rgb = GREEN
    add_para(doc, "Complete Website Documentation & Visual Feature Guide", bold=True)
    add_para(doc, f"Version {VERSION} · Last reviewed {TODAY}")
    add_para(
        doc,
        "This document is generated from the live Finovate Django implementation. "
        "Screenshots use synthetic demo data. Investing features are simulated.",
    )
    doc.add_page_break()

    # TOC-like
    add_heading(doc, "Contents", 1)
    add_bullets(
        doc,
        [
            "1. Product overview & architecture",
            "2. Getting started",
            "3. Feature inventory & routes",
            "4. Authentication & account management",
            "5. Core features (visual walkthroughs)",
            "6. Workflows & business logic",
            "7. Data model, API, integrations",
            "8. Security, deployment, testing",
            "9. Troubleshooting & glossary",
            "10. Coverage matrix & audit notes",
        ],
    )
    doc.add_page_break()

    add_heading(doc, "1. Product overview", 1)
    add_para(
        doc,
        "Finovate is a personal-finance web app: track spending, plan budgets, "
        "pursue savings goals, receive a recommendation-only savings plan, and "
        "preview simulated round-up/ETF investing. Stack: Django 6.1, server-rendered "
        "templates, Decimal money math, SQLite or Postgres, Auth0 and/or Supabase Auth, "
        "optional Twelve Data and Photon Spectrum.",
    )
    add_heading(doc, "System context", 2)
    add_bullets(
        doc,
        [
            "Browser → Django apps → ORM → SQLite/Postgres",
            "Auth0/Supabase for identity (passwords never stored in Finovate)",
            "Twelve Data for optional live demo quotes",
            "Spectrum webhook inbound; Django does not send iMessage",
        ],
    )
    add_heading(doc, "Apps", 2)
    add_bullets(
        doc,
        [
            "core — landing, dashboard, insights, seeds",
            "accounts — profile, auth, onboarding, settings, /api/me/",
            "transactions — accounts, categories, transactions, statements",
            "budgets — periods, limits, income, recurring expenses, txn UI",
            "savings — goals, pools, plan engine, debts",
            "investing — simulated portfolio tools + market data",
            "imports — CSV wizard",
            "assistant — Spectrum intents/webhook",
        ],
    )
    doc.add_page_break()

    add_heading(doc, "2. Getting started", 1)
    add_para(doc, "Local setup:", bold=True)
    add_bullets(
        doc,
        [
            "python -m venv .venv && source .venv/bin/activate",
            "pip install -r requirements.txt",
            "python manage.py migrate",
            "python manage.py seed_demo",
            "python manage.py runserver → http://127.0.0.1:8000/",
        ],
    )
    add_para(
        doc,
        "Website signup/login require Auth0 and/or Supabase configuration. "
        "seed_demo creates a fixture user for management commands and documentation capture.",
    )
    add_para(doc, "Copy .env.example to .env for secrets — never commit real values.", bold=True)
    doc.add_page_break()

    add_heading(doc, "3. Feature inventory (high level)", 1)
    add_bullets(
        doc,
        [
            "Public: landing, login, signup",
            "App: dashboard, insights, transactions, statements, imports",
            "Budgets: periods, limits, income, expenses, categories",
            "Savings: index, pools, goals, plan, debts",
            "Investing: portfolio, round-ups, surplus, demo AI, explore, historical demo",
            "Settings: theme, notifications, messaging link, placeholders",
            "APIs: GET /api/me/, POST /assistant/photon-webhook/, live-quotes JSON",
        ],
    )
    add_para(doc, "Full path listing lives in docs/feature-inventory.md.")
    doc.add_page_break()

    add_heading(doc, "4. Authentication & accounts", 1)
    try_picture(doc, ANN / "authentication-login-annotated.png", 5.8)
    try_picture(doc, RAW / "authentication-signup.png", 5.5)
    add_bullets(
        doc,
        [
            "Auth0: Universal Login; profile keyed by auth0_sub; /callback",
            "Supabase: email/password over HTTPS; profile keyed by supabase_uid",
            "Verify gate for unverified Auth0 emails",
            "Onboarding → optional CSV/manual setup → preferences",
            "Rate limits: signup 10/10m, login 20/10m, resend 3/15m",
            "Logout: Django session clear + Auth0 federated logout when applicable",
        ],
    )
    doc.add_page_break()

    # Visual feature chapters
    chapters = [
        (
            "5. Dashboard",
            "Home of recorded cash, monthly flows, chart, recent transactions, and savings opportunity.",
            [
                ANN / "dashboard-overview-annotated.png",
                RAW / "dashboard-overview.png",
                MOB / "mobile-dashboard.png",
            ],
            [
                "Cards: total cash, monthly income, monthly spending, potential savings",
                "Chart ranges 1W–1Y; simulated portfolio called out separately",
                "CTA into savings plan; recent txn jump list",
            ],
        ),
        (
            "6. Transactions & statements",
            "Manual CRUD, search, CSV export; fictional statement reconciliation demo.",
            [
                ANN / "transactions-list-annotated.png",
                RAW / "transactions-add.png",
                RAW / "statement-detail.png",
            ],
            [
                "Amounts stored positive; kind supplies sign",
                "Decimal ROUND_HALF_UP to 2 places",
                "Peachtree seed demonstrates ledger reconciliation",
            ],
        ),
        (
            "7. CSV import",
            "Upload → account → preview → commit. Max 2MB CSV. No fabricated rows.",
            [ANN / "imports-upload-annotated.png", RAW / "imports-history.png"],
            [
                "ImportBatch tracks status and raw bytes",
                "Dedupe via fingerprints/external_id",
                "Unreadable files guide users to manual setup",
            ],
        ),
        (
            "8. Budgets",
            "Live progress from transactions; income/expense plan estimates.",
            [RAW / "budgets-index.png"],
            [
                "spent = sum(expenses in period/category)",
                "warning/over from percent_used vs threshold",
                "Cadence factors weekly 4.33 / biweekly 2.17",
            ],
        ),
        (
            "9. Savings, pools, goals, plan",
            "Goals with milestones; system pools; recommendation-only plan engine.",
            [
                RAW / "savings-index.png",
                ANN / "savings-plan-annotated.png",
                RAW / "savings-goals.png",
                RAW / "savings-debts.png",
            ],
            [
                "Scenarios: stabilization, build_reserve, debt_focus, stable, strong",
                "Safe-to-save = max(0, surplus − buffer)",
                "Never moves money or places trades",
            ],
        ),
        (
            "10. Simulated investing",
            "Round-ups, surplus proposals, demo AI, explore — all simulated.",
            [
                ANN / "investing-index-annotated.png",
                RAW / "investing-market-data.png",
                RAW / "investing-historical-demo.png",
                MOB / "mobile-investing.png",
            ],
            [
                "Round-up example: $4.35 → $0.65 to nearest $1",
                "Sweep respects pause, cap, cash reserve",
                "Twelve Data optional for live quotes",
            ],
        ),
        (
            "11. Insights & settings",
            "Analytics visuals; preferences, theme, messaging link, placeholders.",
            [
                RAW / "insights-overview.png",
                ANN / "settings-profile-annotated.png",
                RAW / "theme-dark-mode.png",
                MOB / "mobile-settings.png",
            ],
            [
                "Theme light/dark/system via data-theme",
                "Notification events recorded only",
                "Connected accounts are placeholders",
            ],
        ),
    ]

    for title_t, blurb, images, bullets in chapters:
        add_heading(doc, title_t, 1)
        add_para(doc, blurb)
        for img in images:
            try_picture(doc, img, 6.0 if "mobile" not in img.name else 2.8)
        add_bullets(doc, bullets)
        doc.add_page_break()

    add_heading(doc, "12. Key workflows", 1)
    add_heading(doc, "Registration", 2)
    add_bullets(
        doc,
        [
            "Signup → Auth0 or Supabase → profile create by subject/uid",
            "Verify if needed → onboarding → preferences → dashboard",
        ],
    )
    add_heading(doc, "Round-up sweep", 2)
    add_bullets(
        doc,
        [
            "Enable settings → expenses create pending ledger → sweep → simulated portfolio update",
        ],
    )
    add_heading(doc, "Savings plan refresh", 2)
    add_bullets(
        doc,
        [
            "analyze_finances → choose_scenario → build_recommendation → persist if material change",
        ],
    )
    doc.add_page_break()

    add_heading(doc, "13. Data, API, integrations", 1)
    add_para(doc, "Primary models: Profile, Transaction, Category, Account, Budget*, Goal, Contribution, DebtObligation, SavingsPlan*, ImportBatch, RoundUp*, SimulatedOrder, MarketPrice, assistant models.", bold=False)
    add_para(doc, "JSON APIs: GET /api/me/; POST /assistant/photon-webhook/; investing live-quotes endpoints. Other mutations are CSRF-protected form posts.")
    add_bullets(
        doc,
        [
            "Auth0 / Supabase Auth / Postgres",
            "Twelve Data (optional)",
            "Photon Spectrum inbound",
            "Bank/Brokerage keys: placeholders only",
        ],
    )
    doc.add_page_break()

    add_heading(doc, "14. Security, deployment, testing", 1)
    add_bullets(
        doc,
        [
            "Security: IdP auth, CSRF, user-scoped queries, rate limits, webhook signatures, JWT for /api/me/",
            "Deploy: Vercel Django (vercel.json), migrate on boot, prefer Postgres",
            "Test: python manage.py test — broad unit/integration coverage across apps",
        ],
    )
    doc.add_page_break()

    add_heading(doc, "15. Troubleshooting (selected)", 1)
    add_bullets(
        doc,
        [
            "Login unavailable → configure Auth0/Supabase",
            "Callback errors → APP_BASE_URL / Auth0 URL mismatch",
            "Vercel data loss → use Postgres not /tmp SQLite",
            "Twelve Data warnings → key/credits",
            "Investing not in real brokerage → expected (simulated)",
        ],
    )
    doc.add_page_break()

    add_heading(doc, "16. Coverage & audit summary", 1)
    add_para(
        doc,
        "All major product surfaces listed in docs/coverage-matrix.md have Markdown specs and "
        "screenshot coverage except webhook-only assistant (no page) and external Auth0 hosted UI. "
        "Discrepancies (README test count/stack wording, undelivered notifications, echo Spectrum sidecar, "
        "client-only review modal) are reported in docs/audit-report.md without changing production behavior.",
    )
    add_para(doc, "Definition of done: a new developer can set up, trace architecture/processes, test, and troubleshoot; an end user can operate major features via instructions and visuals.", bold=True)

    out = DELIV / "Finovate-Complete-Documentation.docx"
    doc.save(out)
    art = ART / "Finovate-Complete-Documentation.docx"
    art.write_bytes(out.read_bytes())
    print("docx", out, out.stat().st_size)
    return out


def _set_run_font(run, size=18, bold=False, color=None):
    run.font.size = PPt(size)
    run.font.bold = bold
    run.font.name = "Calibri"
    if color is not None:
        run.font.color.rgb = color


def add_slide(prs, title, bullets=None, image=None, notes=None, subtitle=None):
    if image and image.exists():
        layout = prs.slide_layouts[6]  # blank
        slide = prs.slides.add_slide(layout)
        # title bar
        box = slide.shapes.add_textbox(PInches(0.4), PInches(0.25), PInches(9.2), PInches(0.6))
        tf = box.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = title
        _set_run_font(run, 26, True, PptRGB(0x0F, 0x2A, 0x22))
        # image
        max_h = 5.8
        max_w = 9.2
        slide.shapes.add_picture(str(image), PInches(0.4), PInches(0.95), height=PInches(max_h))
        if bullets:
            # small caption under isn't needed; put bullets on next slide if image-heavy
            pass
    else:
        layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = title
        body = slide.shapes.placeholders[1].text_frame
        body.clear()
        if subtitle:
            p = body.paragraphs[0]
            run = p.add_run()
            run.text = subtitle
            _set_run_font(run, 16, False, PptRGB(0x33, 0x33, 0x33))
            first = False
        else:
            first = True
        for i, b in enumerate(bullets or []):
            if first and i == 0 and not subtitle:
                p = body.paragraphs[0]
            else:
                p = body.add_paragraph()
            p.text = b
            p.level = 0
            for run in p.runs:
                _set_run_font(run, 18, False, PptRGB(0x22, 0x22, 0x22))
            first = False
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def generate_pptx() -> Path:
    prs = Presentation()
    prs.slide_width = PInches(10)
    prs.slide_height = PInches(7.5)

    # Title
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # background shape
    shape = slide.shapes.add_shape(1, PInches(0), PInches(0), PInches(10), PInches(7.5))  # rect
    shape.fill.solid()
    shape.fill.fore_color.rgb = PptRGB(0x0F, 0x2A, 0x22)
    shape.line.fill.background()
    tb = slide.shapes.add_textbox(PInches(0.8), PInches(2.2), PInches(8.4), PInches(2))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Finovate"
    _set_run_font(r, 48, True, PptRGB(0xFF, 0xFF, 0xFF))
    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = "Visual Feature Guide & Architecture Walkthrough"
    _set_run_font(r2, 22, False, PptRGB(0xB8, 0xE0, 0xC8))
    p3 = tf.add_paragraph()
    r3 = p3.add_run()
    r3.text = f"Version {VERSION}  ·  {TODAY}  ·  Synthetic demo screenshots"
    _set_run_font(r3, 14, False, PptRGB(0x9A, 0xB5, 0xA8))

    add_slide(
        prs,
        "Agenda",
        [
            "Product purpose & honest boundaries",
            "Architecture & data flow",
            "Authentication & onboarding",
            "Feature visual walkthroughs",
            "Savings plan & simulated investing logic",
            "API, integrations, deployment",
            "Testing, coverage, known gaps",
        ],
    )

    add_slide(
        prs,
        "What Finovate is",
        [
            "Personal finance tracker + planner (Django 6.1)",
            "Real: transactions, budgets, goals, CSV import, plan math",
            "Simulated: investing, round-up sweeps, demo AI orders",
            "Identity: Auth0 and/or Supabase Auth (no local password site login when IdPs expected)",
            "Data: SQLite locally or Postgres (Supabase)",
        ],
        notes="Emphasize simulated investing disclaimers in every demo.",
    )

    add_slide(
        prs,
        "Architecture at a glance",
        [
            "Thin views → services/plan_engine/pipeline → ORM",
            "Decimal money helpers shared in core.services",
            "Per-user ownership on every query",
            "Vercel WSGI entry; migrate-on-boot; optional seed",
            "External: Auth0, Supabase, Twelve Data, Spectrum webhook",
        ],
    )

    # Landing visual
    add_slide(prs, "Landing page", image=ANN / "landing-annotated.png")
    add_slide(
        prs,
        "Landing — callouts",
        [
            "1 Brand / product name as hero signal",
            "2 Sign in / Sign up entry points",
            "3 Value proposition headline",
            "4 Primary CTA into registration",
            "Public route: /",
        ],
    )

    add_slide(prs, "Login", image=ANN / "authentication-login-annotated.png")
    add_slide(prs, "Sign up", image=RAW / "authentication-signup.png")
    add_slide(
        prs,
        "Auth lifecycle",
        [
            "Signup/login via Auth0 Universal Login or Supabase email",
            "Profile keyed by auth0_sub or supabase_uid (never email merge for Auth0)",
            "Verify gate → onboarding → preferences → dashboard",
            "Rate limits protect signup/login/resend",
            "GET /api/me/ accepts Auth0 Bearer access tokens",
        ],
    )

    add_slide(prs, "Onboarding", image=RAW / "onboarding-start.png")
    add_slide(
        prs,
        "Onboarding paths",
        [
            "Import CSV bank statement (real commit after preview)",
            "Connect bank — coming soon placeholder",
            "Manual setup estimates → budget plan summary",
            "Preferences set onboarding_completed",
        ],
    )

    add_slide(prs, "Dashboard overview", image=ANN / "dashboard-overview-annotated.png")
    add_slide(
        prs,
        "Dashboard elements",
        [
            "1 Sidebar navigation across product areas",
            "2 Account header, notifications affordance, logout",
            "3 Cash & monthly summary cards",
            "4 Cash balance chart (range controls)",
            "5 Savings opportunity / plan CTA",
            "6 Recent transactions",
        ],
    )
    add_slide(prs, "Dashboard (mobile)", image=MOB / "mobile-dashboard.png")

    add_slide(prs, "Transactions", image=ANN / "transactions-list-annotated.png")
    add_slide(prs, "Add transaction", image=RAW / "transactions-add.png")
    add_slide(prs, "Statement detail", image=RAW / "statement-detail.png")
    add_slide(
        prs,
        "Transactions — logic",
        [
            "CRUD under /budgets/transactions/",
            "CSV export endpoint",
            "Positive amount + kind; Decimal quantization",
            "Idempotent imports via external_id",
            "Statements reconcile computed vs stated balances",
        ],
    )

    add_slide(prs, "CSV import upload", image=ANN / "imports-upload-annotated.png")
    add_slide(
        prs,
        "Import wizard",
        [
            "Upload (≤2MB CSV) → choose/create account → preview → commit",
            "ImportBatch stores raw file + mapping + counters",
            "Duplicates skipped; unreadable → manual setup path",
            "History/retry/delete at /imports/",
        ],
    )

    add_slide(prs, "Budgets", image=RAW / "budgets-index.png")
    add_slide(
        prs,
        "Budget math",
        [
            "Progress spent is live-summed from transactions (never stored)",
            "Warning when percent_used ≥ threshold and < 100; over at ≥ 100",
            "Plan cadence factors: weekly×4.33, biweekly×2.17",
            "Suggested savings heuristic uses remaining goals ÷ 12",
        ],
    )

    add_slide(prs, "Savings home", image=RAW / "savings-index.png")
    add_slide(prs, "Financial goals", image=RAW / "savings-goals.png")
    add_slide(prs, "Savings plan", image=ANN / "savings-plan-annotated.png")
    add_slide(prs, "Debts (plan inputs)", image=RAW / "savings-debts.png")
    add_slide(
        prs,
        "Savings plan engine",
        [
            "Recommendation only — never transfers funds",
            "Safe-to-save = max(0, surplus − buffer)",
            "Scenarios: stabilization → build reserve → debt focus → stable → strong",
            "Default pools: Rainy Day, Investment, Regular Savings",
            "Revisions stored when material changes occur",
        ],
    )

    add_slide(prs, "Investing (simulated)", image=ANN / "investing-index-annotated.png")
    add_slide(prs, "Market data explore", image=RAW / "investing-market-data.png")
    add_slide(prs, "Historical demo", image=RAW / "investing-historical-demo.png")
    add_slide(
        prs,
        "Investing rules",
        [
            "No brokerage connected — UI disclaimers required",
            "Round-up: $4.35 with $1 increment → $0.65",
            "Sweep enforces pause, contribution cap, cash reserve",
            "Surplus propose → approve/reject SimulatedOrder",
            "Demo AI is rule-based; Twelve Data optional",
        ],
    )

    add_slide(prs, "Insights", image=RAW / "insights-overview.png")
    add_slide(prs, "Settings", image=ANN / "settings-profile-annotated.png")
    add_slide(prs, "Dark theme", image=RAW / "theme-dark-mode.png")
    add_slide(
        prs,
        "Interface behaviors",
        [
            "Themes: light / dark / system via data-theme",
            "Mobile nav toggle; desktop sidebar",
            "Live quote polling JS (~5 minutes)",
            "Review modal is client-only (does not persist plan)",
            "Notifications recorded, not emailed/pushed",
        ],
    )

    add_slide(
        prs,
        "API surface",
        [
            "GET /api/me/ — Auth0 Bearer token → identity + txn count",
            "POST /assistant/photon-webhook/ — signed Spectrum inbound",
            "GET /investing/live-quotes/ — session JSON holdings quotes",
            "GET /investing/market-data/live-quotes/ — session market quotes",
            "All other mutations: CSRF-protected HTML form posts",
        ],
    )
    add_slide(
        prs,
        "Integrations",
        [
            "Auth0 — primary OIDC identity",
            "Supabase Auth — optional email/password",
            "Postgres via DATABASE_URL / SUPABASE_DB_URL",
            "Twelve Data — optional live market data",
            "Photon Spectrum — inbound messaging only from Django’s POV",
            "Bank/Brokerage API keys — placeholders (unused)",
        ],
    )
    add_slide(
        prs,
        "Security controls",
        [
            "IdP authentication; passwords not stored in Finovate",
            "User-scoped ORM filters on every data path",
            "CSRF on forms; session auth for pages",
            "IP rate limits on signup/login/resend",
            "Webhook HMAC / shared secret verification",
            "Env-based secrets; never commit real values",
        ],
    )
    add_slide(
        prs,
        "Deployment & operations",
        [
            "Vercel Django framework entry: finovate/wsgi.py",
            "maxDuration 60s; migrate on boot when VERCEL set",
            "Prefer Postgres — /tmp SQLite is ephemeral",
            "Configure APP_BASE_URL + Auth0 callbacks",
            "python manage.py test for regression suite",
        ],
    )
    add_slide(
        prs,
        "Troubleshooting highlights",
        [
            "Login/signup inert → missing Auth0/Supabase env",
            "Callback failures → APP_BASE_URL mismatch",
            "Verify loop → confirm email in Auth0, then refresh",
            "Import errors → size/format; use manual setup",
            "Twelve Data warnings → key or free-tier credits",
            "“Trades” not in brokerage → expected (simulated)",
        ],
    )
    add_slide(
        prs,
        "Documentation deliverables",
        [
            "Markdown portal under /docs (navigable README)",
            "Word: docs/deliverables/Finovate-Complete-Documentation.docx",
            "PowerPoint: docs/deliverables/Finovate-Visual-Feature-Guide.pptx",
            "Screenshots: docs/assets/screenshots/{raw,annotated,mobile}",
            "Coverage matrix + audit report included",
        ],
    )
    add_slide(
        prs,
        "Known gaps / defects reported",
        [
            "Root README test count / Tailwind mentions outdated vs code",
            "Notifications toggles without delivery channel",
            "Spectrum Bun sidecar currently echo-only",
            "Bank connect coming soon; brokerage simulated",
            "No transaction pagination yet",
            "Optional follow-up: validation-error & import-preview shots",
        ],
    )

    # Closing
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_shape(1, PInches(0), PInches(0), PInches(10), PInches(7.5))
    shape.fill.solid()
    shape.fill.fore_color.rgb = PptRGB(0x1F, 0x6B, 0x4A)
    shape.line.fill.background()
    tb = slide.shapes.add_textbox(PInches(0.8), PInches(2.4), PInches(8.4), PInches(2.5))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Thank you"
    _set_run_font(r, 44, True, PptRGB(0xFF, 0xFF, 0xFF))
    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = "See /docs for the full reference. Investing is simulated — no real trades."
    _set_run_font(r2, 16, False, PptRGB(0xE0, 0xF2, 0xE8))

    out = DELIV / "Finovate-Visual-Feature-Guide.pptx"
    prs.save(out)
    art = ART / "Finovate-Visual-Feature-Guide.pptx"
    art.write_bytes(out.read_bytes())
    print("pptx", out, out.stat().st_size)
    return out


def main() -> None:
    generate_markdown()
    generate_docx()
    generate_pptx()
    print("ALL DONE")


if __name__ == "__main__":
    main()